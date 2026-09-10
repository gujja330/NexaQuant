"""FIELD COMPLETENESS · every column the operator reads must be populated.

> "many gaps, fields missing like confident, stop or whatever make
>  thorough check and its ur responsibility to check"

WHY THIS IS A GATE AND NOT A PROMISE
-------------------------------------
"I will check the fields" is worth nothing on the day I forget. A row with
a blank Stop looks almost identical to a row with a real one, and the
operator only discovers the difference when they act on it. So the check
runs every time, names the exact column and ticker, and blocks delivery -
the same standing the freshness and conservation gates have.

WHAT COUNTS AS MISSING
----------------------
Not just None. `""`, `"None"`, `"nan"`, `"—"` and `"UNAVAILABLE"` are all
ways a pipeline says "I had nothing" while still filling the cell. Each is
treated as absent, because to a reader they are.

The em-dash is the one deliberate exception, and only where the delivery
contract already governs it: a Path-A holding legitimately renders `—`
rather than a fabricated number. That exemption is declared per field
below, never inferred.

REQUIRED vs EXPECTED
--------------------
REQUIRED fields block. A NEW recommendation without an Entry Price or a
Stop is not shippable - it cannot be acted on.

EXPECTED fields warn. Confidence on an R1 advisory row is often genuinely
absent upstream, and blocking on it would stop delivery for something the
operator does not trade from.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from backend.pipeline.contract.run_context import RunContext, StageContract

SCHEMA_VERSION = "aegis.pipeline.contract.completeness.v1"

STAGE = "FIELDS_COMPLETE"

# Values that occupy a cell while meaning "nothing".
EMPTY_TOKENS = {"", "none", "nan", "null", "n/a", "na", "-",
                "unavailable", "not_available"}

# Fields an actionable row cannot ship without.
REQUIRED_BY_ACTION = {
    "NEW":      ("Entry Price", "Last Price", "Stop", "Confidence %"),
    "ACTIVE+":  ("Entry Price", "Last Price", "Stop", "P&L %"),
    "ACTIVE":   ("Entry Price", "Last Price", "Stop", "P&L %"),
}

# Fields that should be present but whose absence is upstream, not a bug
# worth stopping a delivery for.
EXPECTED = ("Target", "Stop Distance %", "Stop State", "Reason",
            "Sector", "Market Data As-of", "Admission Status")

# Fields where the delivery contract explicitly governs an em-dash.
#
# Confidence was in this set and that was the hole: the twelve USA NEW
# rows of 2026-09-09 rendered Confidence as "—" and this gate passed them,
# while the SSOT held 0.5905 all along. An em-dash is legitimate on a
# Path-A HOLDING whose confidence genuinely predates the field; it is
# never legitimate on a NEW recommendation, which cannot be judged
# without one.
DASH_ALLOWED = {"Target", "Stop Distance %", "Size"}

# Explicit "we have no value" strings. They are honest for a HOLDING
# whose score predates the field, and they must NEVER satisfy a NEW
# recommendation - a name being proposed today cannot be judged without
# a confidence, and a sentence explaining its absence is not a value.
UNAVAILABLE_TOKENS = {"historical score unavailable",
                      "market_cap_unavailable"}

# Per-action exemptions · an em-dash here is governed, not a gap.
DASH_ALLOWED_BY_ACTION = {
    "ACTIVE": {"Confidence %"},      # legacy holdings predate the field
    "ACTIVE+": {"Confidence %"},
}

# R1 is retired advisory · it is not traded from, so it is not held to the
# actionable-row standard.
ADVISORY_ENGINES = {"R1"}


def _blank(v, field: str, action: str = "") -> bool:
    if v is None:
        return True
    s = str(v).strip()
    if s == "—":                       # em-dash
        if field in DASH_ALLOWED:
            return False
        return field not in DASH_ALLOWED_BY_ACTION.get(
            str(action).upper(), set())
    if s.lower() in UNAVAILABLE_TOKENS:
        # Governed absence · acceptable on a holding, never on a NEW row.
        return str(action).upper() == "NEW"
    return s.lower() in EMPTY_TOKENS


# Rounding slack for the stop-distance identity. Displayed values are
# rounded to 2dp, so anything inside this is presentation, not disagreement.
STOP_DISTANCE_TOLERANCE_PCT = 0.05


def stop_distance_invariant(root: Path, market: str) -> dict:
    """Stop, price and distance must be able to agree ARITHMETICALLY.

        distance = |price - stop| / price * 100

    Displayed together, these three numbers are a claim about downside.
    If they cannot be reconciled the reader cannot tell which is wrong,
    and a stop is the one number on the sheet someone acts on in a hurry.

    Nothing is repaired here. A row that cannot agree FAILS DELIVERY, so
    a silently corrected display can never ship.

    A `target` at or below the current price is also incoherent - it is a
    target already passed - and is reported so it can be withheld rather
    than shown.
    """
    import json
    p = (Path(root) / "reports" / "context"
         / ("canonical_lifecycle_%s.json" % market.lower()))
    if not p.exists():
        return {"error": "lifecycle dataset missing"}
    d = json.loads(p.read_text(encoding="utf-8"))
    rows = d.get("current") or []
    out = {"rows_checked": len(rows), "rows_with_stop": 0, "mismatches": [],
           "malformed_stops": [], "missing_stops": [],
           "incoherent_targets": []}
    for r in rows:
        tk = str(r.get("ticker"))
        stop, price = r.get("stop"), r.get("current_price")
        if not isinstance(stop, (int, float)):
            out["missing_stops"].append(tk)
            continue
        if stop <= 0 or not isinstance(price, (int, float)) or price <= 0:
            out["malformed_stops"].append(
                {"ticker": tk, "stop": stop, "price": price})
            continue
        out["rows_with_stop"] += 1
        expected = abs(price - stop) / price * 100
        shown = r.get("dist_to_stop_pct")
        if isinstance(shown, (int, float)) and                 abs(expected - shown) > STOP_DISTANCE_TOLERANCE_PCT:
            out["mismatches"].append({
                "ticker": tk, "engine": r.get("engine"), "price": price,
                "stop": stop, "displayed_distance": shown,
                "expected_distance": round(expected, 4)})
        t = r.get("target")
        if isinstance(t, (int, float)) and t <= price:
            out["incoherent_targets"].append({
                "ticker": tk, "engine": r.get("engine"), "target": t,
                "current_price": price})
    out["pass"] = not (out["mismatches"] or out["malformed_stops"])
    return out


def scan_workbook(root: Path, market: str) -> dict:
    """Read the RENDERED sheet · what the operator will actually see."""
    from openpyxl import load_workbook

    from backend.delivery.xlsx_contract import resolve_exit_history_sheet
    p = root / "reports" / "telegram" / ("aegis_history_%s.xlsx" % market)
    if not p.exists():
        return {"error": "workbook missing: %s" % p.name}
    wb = load_workbook(p, read_only=True)
    out = {"sheets": list(wb.sheetnames), "required_gaps": [],
           "expected_gaps": [], "rows_checked": 0}

    if "CURRENT" in wb.sheetnames:
        ws = wb["CURRENT"]
        hdr = [str(c.value).strip() if c.value else "" for c in ws[7]]
        idx = {h: i for i, h in enumerate(hdr) if h}
        for r in range(8, ws.max_row + 1):
            tk = ws.cell(r, idx.get("Ticker", 1) + 1).value
            if not tk or str(tk).startswith("─"):
                break
            out["rows_checked"] += 1
            row = {h: ws.cell(r, i + 1).value for h, i in idx.items()}
            action = str(row.get("Action") or "").upper()
            engine = str(row.get("Engine") or "").upper()
            req = REQUIRED_BY_ACTION.get(action, ())
            for f in req:
                if f not in idx:
                    out["required_gaps"].append(
                        {"ticker": str(tk), "field": f, "why": "column absent"})
                elif _blank(row.get(f), f, action):
                    # An advisory row is not actionable · record, do not block.
                    bucket = ("expected_gaps" if engine in ADVISORY_ENGINES
                              else "required_gaps")
                    out[bucket].append({"ticker": str(tk), "action": action,
                                        "engine": engine, "field": f,
                                        "value": str(row.get(f))})
            for f in EXPECTED:
                if f in idx and _blank(row.get(f), f, action):
                    out["expected_gaps"].append(
                        {"ticker": str(tk), "field": f,
                         "value": str(row.get(f))})

    if "EXIT HISTORY" in wb.sheetnames:
        ws = wb["EXIT HISTORY"]
        hdr = [str(c.value).strip() if c.value else "" for c in ws[4]]
        idx = {h: i for i, h in enumerate(hdr) if h}
        n_ex, sect = 0, 0
        for r in range(5, ws.max_row + 1):
            if ws.cell(r, 1).value in (None, ""):
                break
            n_ex += 1
            if "Sector" in idx and not _blank(
                    ws.cell(r, idx["Sector"] + 1).value, "Sector"):
                sect += 1
        out["exit_rows"] = n_ex
        out["exit_sector_populated"] = sect
    wb.close()
    return out


def check_fields_complete(ctx: RunContext) -> StageContract:
    import time
    t0 = time.time()
    c = ctx.stage(STAGE)
    root, m = Path(ctx.root), ctx.market

    why = ctx.upstream_ok("LIFECYCLE_READY")
    if why:
        c.block("UPSTREAM_NOT_READY", why)
        c.elapsed_s = round(time.time() - t0, 2)
        return ctx.record(c)

    res = scan_workbook(root, m)
    if res.get("error"):
        c.block("WORKBOOK_MISSING", res["error"])
        c.elapsed_s = round(time.time() - t0, 2)
        return ctx.record(c)

    c.row_count = res["rows_checked"]
    c.detail["expected_gaps"] = res["expected_gaps"][:20]
    c.detail["n_expected_gaps"] = len(res["expected_gaps"])
    c.detail["exit_rows"] = res.get("exit_rows")
    c.detail["exit_sector_populated"] = res.get("exit_sector_populated")
    req = res["required_gaps"]
    c.detail["n_required_gaps"] = len(req)
    c.detail["required_gaps"] = req[:20]
    c.missing_count = len(req)

    # ── stop / price / distance must agree · CEO 2026-09-09 ───────────
    inv = stop_distance_invariant(root, m)
    c.detail["stop_distance"] = {
        k: inv.get(k) for k in
        ("rows_checked", "rows_with_stop", "mismatches", "malformed_stops",
         "missing_stops", "incoherent_targets")}
    if inv.get("mismatches"):
        c.block("STOP_DISTANCE_MISMATCH",
                "%d row(s) where stop, price and distance cannot agree: %s · "
                "a displayed stop that does not reconcile is not shippable"
                % (len(inv["mismatches"]),
                   ", ".join(x["ticker"] for x in inv["mismatches"][:5])))
    if inv.get("malformed_stops"):
        c.block("STOP_MALFORMED",
                "%d row(s) carry a non-positive stop or price: %s"
                % (len(inv["malformed_stops"]),
                   ", ".join(x["ticker"] for x in inv["malformed_stops"][:5])))

    if req:
        by_field = {}
        for g in req:
            by_field.setdefault(g["field"], []).append(g["ticker"])
        summary = "; ".join("%s missing on %d row(s): %s"
                            % (f, len(t), ", ".join(sorted(set(t))[:5]))
                            for f, t in sorted(by_field.items()))
        c.block("REQUIRED_FIELD_MISSING",
                "an actionable row cannot ship without these · " + summary)

    c.elapsed_s = round(time.time() - t0, 2)
    return ctx.record(c.ok())
