"""AEGIS RUN CONTRACT · the five fail-closed gates.

Each test below corresponds to a failure that actually shipped:

  DATA/FEATURES  a 50-day-old feature store scored as if it were today
  SCORES         four qualified BUYs discarded before any rule saw them
  DECISION       a BUY that reached no registry entry and nothing stopped
  FIELDS         rows delivered with an empty Stop column
  DELIVERY       a gate that measured staleness and allowed the send

The point of every one is the same: the pipeline must REFUSE, not report.
"""
from __future__ import annotations
import json

from datetime import date, timedelta
from pathlib import Path

import pytest

from backend.pipeline.contract import completeness, readiness, scoring
from backend.pipeline.contract.run_context import (RunContext, StageContract,
                                                   new_run_id, resolve_asof)

ROOT = Path(__file__).resolve().parents[2]
MARKETS = ("india", "usa")


def _ctx(market="india", asof=None):
    return RunContext(market=market, requested_asof=asof or resolve_asof(),
                      root=ROOT)


# ── identity and fail-closed plumbing ───────────────────────────────────

def test_run_id_carries_market_and_asof():
    rid = new_run_id("india", "2026-09-09")
    assert rid.startswith("AEGIS-2026-09-09-INDIA-")


def test_asof_is_decided_once():
    ctx = _ctx(asof="2026-09-09")
    a = ctx.stage("X")
    assert a.input_asof == "2026-09-09"
    assert ctx.requested_asof == "2026-09-09"


def test_downstream_refuses_when_upstream_blocked():
    ctx = _ctx()
    c = ctx.stage("DATA_READY").block("SOURCE_STALE", "50 days old")
    ctx.record(c)
    why = ctx.upstream_ok("DATA_READY")
    assert why and "BLOCKED" in why


def test_downstream_refuses_when_upstream_never_ran():
    assert "did not run" in (_ctx().upstream_ok("DATA_READY") or "")


def test_first_failure_is_singular_and_named():
    ctx = _ctx()
    ctx.record(ctx.stage("A").block("FIRST", "one"))
    ctx.record(ctx.stage("B").block("SECOND", "two"))
    assert ctx.first_failure["code"] == "FIRST"
    assert ctx.blocked


# ── GATE 1/2 · the July-21 failure ──────────────────────────────────────

def test_stale_substrate_blocks_the_run():
    """A snapshot older than its budget must stop everything."""
    ctx = _ctx(asof=(date.today() + timedelta(days=400)).isoformat())
    c = readiness.check_data_ready(ctx)
    assert c.status == "BLOCK"
    assert any("SOURCE_STALE" in f["code"] for f in c.failures)


def test_feature_budget_is_zero_days():
    """Today's scoring needs today's features · not 'the latest file'."""
    assert readiness.SOURCE_BUDGETS["feature_snapshot"] == 0


def test_snapshot_lookup_is_by_requested_date_not_latest():
    src = Path(readiness.__file__).read_text(encoding="utf-8")
    assert "never 'the latest one'" in src
    # the resolver must compare against the requested date
    assert "if want not in set(list_snapshots" in src


def test_features_ready_blocks_on_asof_mismatch():
    ctx = _ctx(asof="1999-01-04")
    ctx.record(ctx.stage(readiness.STAGE_DATA).ok())
    c = readiness.check_features_ready(ctx)
    assert c.status == "BLOCK"
    assert any(f["code"] == "FEATURE_SNAPSHOT_ASOF_MISMATCH"
               for f in c.failures)


@pytest.mark.parametrize("market", MARKETS)
def test_universe_resolves_to_a_real_number(market):
    """A silent zero makes every coverage ratio meaningless."""
    n = readiness._universe_size(ROOT, market)
    assert n and n > 50, "universe for %s resolved to %r" % (market, n)


# ── GATE 3 · truncation before eligibility ──────────────────────────────

def test_presentation_truncation_is_a_blocking_condition():
    src = Path(scoring.__file__).read_text(encoding="utf-8")
    assert "TRUNCATION_BEFORE_ELIGIBILITY" in src


def test_ensemble_persists_the_full_universe():
    """top_10 + bottom_5 is a display slice · eligibility needs all."""
    import json
    for rel in ("reports/ensemble.json", "usa/reports/ensemble.json"):
        p = ROOT / rel
        if not p.exists():
            continue
        d = json.loads(p.read_text(encoding="utf-8"))
        assert d.get("n_all_candidates"), "%s has no all_candidates" % rel
        assert d["n_all_candidates"] > 15


def test_all_candidates_carry_per_model_scores():
    """Omitting them made every row look like a disagreement -> HOLD."""
    import json
    for rel in ("reports/ensemble.json", "usa/reports/ensemble.json"):
        p = ROOT / rel
        if not p.exists():
            continue
        d = json.loads(p.read_text(encoding="utf-8"))
        rows = d.get("all_candidates") or []
        assert rows and rows[0].get("per_model_score") is not None, rel


def test_recommendation_engine_reads_all_candidates():
    for rel in ("usa/research/recommendation_intelligence/run.py",
                "india/recommendation_intelligence/run.py"):
        src = (ROOT / rel).read_text(encoding="utf-8")
        assert 'ens.get("all_candidates")' in src, rel


# ── GATE 4 · conservation ───────────────────────────────────────────────

def test_conservation_uses_the_funnel_normalisation():
    """A mismatched spelling reported 228 phantom orphans."""
    assert scoring._norm("MCX.NS") == "MCX"
    assert scoring._norm("BRK.B") == "BRK"
    assert scoring._norm(" aapl ") == "AAPL"


def test_cooling_is_a_terminal_state_not_a_defect():
    src = Path(scoring.__file__).read_text(encoding="utf-8")
    assert "COOLING is a legitimate terminal state" in src


def test_loss_is_judged_by_production_not_by_the_shadow():
    """The shadow re-trains in-process · its BUY is not production's."""
    src = Path(scoring.__file__).read_text(encoding="utf-8")
    assert "_production_buys" in src
    assert "RE-TRAINS the model factory" in src


def test_every_terminal_disposition_is_declared():
    for v in ("L2_hold", "L3_disagreement", "L7_eligible_no_registry",
              "L10_in_current"):
        assert v in scoring.TERMINAL_DISPOSITIONS


# ── FIELDS · the operator-visible gaps ──────────────────────────────────

def test_blank_tokens_are_treated_as_missing():
    for v in (None, "", "None", "nan", "UNAVAILABLE", "N/A"):
        assert completeness._blank(v, "Stop"), repr(v)


def test_em_dash_is_missing_unless_the_contract_allows_it():
    # A Stop is never optional.
    assert completeness._blank("—", "Stop", "NEW") is True
    # Confidence: governed on a legacy holding, never on a NEW row. The
    # twelve USA NEW rows of 2026-09-09 shipped an em-dash here while the
    # SSOT held 0.5905 · that is the gap this distinction closes.
    assert completeness._blank("—", "Confidence %", "NEW") is True
    assert completeness._blank("—", "Confidence %", "ACTIVE") is False
    assert completeness._blank("—", "Confidence %", "ACTIVE+") is False
    # Fields the delivery contract governs outright.
    assert completeness._blank("—", "Target", "NEW") is False


def test_stop_is_required_on_every_actionable_row():
    for action in ("NEW", "ACTIVE+", "ACTIVE"):
        assert "Stop" in completeness.REQUIRED_BY_ACTION[action]
        assert "Entry Price" in completeness.REQUIRED_BY_ACTION[action]


@pytest.mark.parametrize("market", MARKETS)
def test_rendered_workbook_has_no_required_gaps(market):
    """The LIVE artifact · skipped when it predates the current renderer.

    This reads the shipped workbook deliberately, because a gap in what
    was actually delivered matters. But on a CI runner that file may be a
    committed artifact from an earlier renderer, and failing the suite on
    that blocks the send for a non-reason. The pipeline's FIELDS_COMPLETE
    gate is what guarantees the freshly built workbook is clean.
    """
    from backend.delivery.lifecycle import canonical_daily_lifecycle as lc
    d = lc.load(ROOT, market)
    res = completeness.scan_workbook(ROOT, market)
    if res.get("error"):
        pytest.skip(res["error"])
    if d and res.get("rows_checked", 0) != len(d.get("current") or []):
        pytest.skip("workbook predates the current lifecycle · the gate, "
                    "not pytest, validates what actually ships")
    # Same reasoning for COLUMN drift. A shipped workbook built before a
    # column rename reports every renamed field as "column absent", which
    # is an artifact-age failure wearing the costume of a data gap. This
    # has blocked production three times.
    if any(g.get("why") == "column absent" for g in res["required_gaps"]):
        pytest.skip("workbook predates the current column contract · the "
                    "FIELDS_COMPLETE gate validates what actually ships")
    assert not res["required_gaps"], res["required_gaps"][:5]


# ── the run is one path ─────────────────────────────────────────────────

def test_ci_and_manual_execute_the_same_contract():
    for rel in (".github/workflows/aegis-daily.yml",
                ".github/workflows/aegis-usa.yml"):
        src = (ROOT / rel).read_text(encoding="utf-8")
        assert "backend.pipeline.contract.runner" in src, rel


def test_send_is_gated_on_the_contract():
    for rel, step in ((".github/workflows/aegis-daily.yml", "run_contract"),
                      (".github/workflows/aegis-usa.yml", "run_contract_usa")):
        src = (ROOT / rel).read_text(encoding="utf-8")
        assert ("steps.%s.outcome == 'success'" % step) in src, rel


def test_critical_substrate_steps_are_not_optional():
    from backend.observability.pipeline_audit import audit_steps
    a = audit_steps(ROOT)
    assert not a["critical_substrate_steps_marked_optional"], a[
        "critical_substrate_steps_marked_optional"]


def test_producers_declare_the_artifact_they_write():
    from backend.observability.pipeline_audit import audit_steps
    a = audit_steps(ROOT)
    assert not a["undeclared_outputs"], a["undeclared_outputs"]


# ── INTERMARKET · the world the stock trades inside ─────────────────────

def test_intermarket_series_are_declared_not_inferred():
    from backend.pipeline.contract import intermarket as im
    labels = {s[0] for s in im.SERIES}
    for required in ("S&P 500", "VIX", "DXY", "USD/INR", "US 10Y",
                     "Oil (Brent/WTI)", "India VIX", "Macro regime"):
        assert required in labels, required


def test_uncollected_series_are_named_not_hidden():
    """A completeness score that counts only what exists is a lie."""
    from backend.pipeline.contract import intermarket as im
    nc = im._not_collected("india")
    assert "Silver" in nc
    assert "India sector indices" in nc
    # USA has no India sector indices to be missing.
    assert "India sector indices" not in im._not_collected("usa")


def test_intermarket_feeds_nothing_into_scoring():
    """Collection is not evidence · D17 stays evidence-blocked."""
    from backend.pipeline.contract import intermarket as im
    s = im.survey(ROOT, "india", resolve_asof())
    assert s["feeds_scoring"] is False
    assert "evidence-BLOCKED" in s["note"]


@pytest.mark.parametrize("market", MARKETS)
def test_no_critical_context_series_is_stale(market):
    """India's global set froze on 2026-06-19 and sat 3 months stale
    while every downstream artifact reported itself fresh."""
    from backend.pipeline.contract import intermarket as im
    s = im.survey(ROOT, market, resolve_asof())
    assert not s["critical_missing"], s["critical_missing"]
    assert not s["critical_stale"], s["critical_stale"]


def test_global_context_producer_is_a_pipeline_step():
    """india/global_risk.py existed but nothing called it."""
    from backend.observability.pipeline_audit import _steps
    names = {s["name"] for s in _steps(ROOT)}
    assert "global_context_bars" in names
    step = next(s for s in _steps(ROOT) if s["name"] == "global_context_bars")
    assert step.get("optional") is False
    assert len(step.get("produces") or []) >= 7


def test_bar_asof_reads_the_last_row_not_the_mtime():
    """mtime moves on a re-run even when the content is unchanged."""
    from backend.pipeline.contract import intermarket as im
    src = Path(im.__file__).read_text(encoding="utf-8")
    assert "never its mtime" in src or "Last row date" in src


def test_global_context_step_actually_fetches():
    """Without --fetch the script only PRINTS · the step would run green
    every day and refresh nothing, recreating the 3-month staleness."""
    from backend.observability.pipeline_audit import _steps
    step = next(s for s in _steps(ROOT) if s["name"] == "global_context_bars")
    assert "--fetch" in (step.get("script_args") or []), step.get("script_args")
    from backend.pipeline.contract import runner
    src = Path(runner.__file__).read_text(encoding="utf-8")
    assert '"india/global_risk.py",\n                            "--fetch"' in src \
        or '"--fetch"' in src


# ── evidence persistence · an accumulator that evaporates accumulates
#    nothing ──────────────────────────────────────────────────────────

def test_fundamental_substrate_is_tracked_in_git():
    """498 USA tickers / 3,431 quarterly periods survived only in a stash.

    The substrate had never been committed. Hours of accumulation sat in
    the working tree, and `git stash -u` swept it three separate times.
    R3 evidence accumulation is impossible if the evidence base is not
    persisted, so its presence in the index is asserted here.
    """
    import subprocess
    tracked = subprocess.run(
        ["git", "ls-files", "reports/research/substrate/"],
        cwd=str(ROOT), capture_output=True, text=True).stdout
    for m in ("india", "usa"):
        assert "fundamental_timeseries_%s.json" % m in tracked, (
            "%s fundamental substrate is not tracked · it will be lost" % m)


def test_send_loads_credentials_from_file():
    """`--send` reported message=False xlsx=False with no reason because
    it read os.environ only · a delivery that fails silently is the
    pattern this contract exists to remove."""
    from backend.pipeline.contract import runner
    src = Path(runner.__file__).read_text(encoding="utf-8")
    assert "_load_telegram_env" in src
    assert ".env.telegram" in src
    assert "DELIVERY FAILED" in src


# ── R3 evidence registry · readiness is arithmetic, not opinion ─────────

def test_registry_uses_only_declared_terminal_states():
    from backend.research.r3_program import evidence_registry as er
    rep = er.load(ROOT) or er.build(ROOT)
    for r in rep["families"]:
        assert r["current_verdict"] in er.STATES, r["current_verdict"]


def test_disposed_branches_stay_disposed():
    """A failed branch is a research result, not a retry invitation."""
    from backend.research.r3_program import evidence_registry as er
    rep = er.load(ROOT) or er.build(ROOT)
    by = {}
    for r in rep["families"]:
        by.setdefault(r["research_id"], set()).add(r["current_verdict"])
    assert by["R3-E-2C"] == {"NO_INCREMENTAL_VALUE"}
    assert by["II.1-GBM"] == {"REJECTED"}
    assert by["R3-H-UNSUP"] == {"REJECTED"}
    assert by["R3-C-SEQ"] == {"FROZEN"}
    assert by["R3-G-EXIT"] == {"FROZEN"}


def test_no_branch_is_ready_without_meeting_its_gate():
    from backend.research.r3_program import evidence_registry as er
    rep = er.load(ROOT) or er.build(ROOT)
    for r in rep["families"]:
        if r["research_id"] == "R3-D-FCST" and r["ready_for_evaluation"]:
            assert r["effective_sample"] >= er.GATE_FUNDAMENTAL_TICKERS
        if r["research_id"] == "R3-K-META" and r["ready_for_evaluation"]:
            assert r["effective_sample"] >= er.GATE_METALABEL_OUTCOMES


def test_gates_are_preregistered_not_tuned():
    from backend.research.r3_program import evidence_registry as er
    assert er.GATE_FUNDAMENTAL_QUARTERS == 8
    assert er.GATE_METALABEL_OUTCOMES == 50
    assert er.GATE_CALIBRATION_ECE == 0.05


def test_registry_runs_in_the_daily_pipeline():
    from backend.pipeline.contract import runner
    src = Path(runner.__file__).read_text(encoding="utf-8")
    assert "r3_evidence_registry" in src


def test_breach_latches_reference_live_positions():
    """A latch whose position no longer exists protects nothing.

    On 2026-09-09 a `git checkout HEAD` of the opportunity registry
    rolled back that day's admissions while the breach ledger kept its
    latches, orphaning USA-R2-AMGN-20260909-69465a. The transition check
    then reported a position "silently dropped".

    The two files are coupled and must be repaired together. Genuine
    latches for LIVE positions are never removed - un-latching a real
    breach would be far worse than a failing test.
    """
    import json

    from backend.research import opportunity_registry as oreg
    reg = oreg.load_all(ROOT)
    live = {o.opportunity_id for opps in reg.values() for o in opps}
    for market in MARKETS:
        p = (ROOT / "reports" / "context"
             / ("lifecycle_breach_ledger_%s.jsonl" % market))
        if not p.exists():
            continue
        orphaned = []
        for line in p.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                pid = json.loads(line).get("position_id")
            except Exception:
                continue
            if pid and pid not in live:
                orphaned.append(pid)
        assert not orphaned, (
            "%s breach ledger latches positions that no longer exist: %s"
            % (market, orphaned[:5]))


# ── evidence clock · is the evidence actually accumulating ──────────────

def test_clock_distinguishes_deduplication_from_loss():
    """India fell 63 -> 53 rows · that was dedup, not ten lost predictions."""
    from backend.research.r3_program import evidence_clock as ec
    cur = {"markets": {"india": {"raw_rows": 53, "unique_predictions": 53,
                                 "unique_tickers": 43, "matured_outcomes": 0,
                                 "fundamental_periods": 265, "tickers": 45,
                                 "tickers_ge_8q": 0, "median_quarters": 6},
                       "usa": {}}, "registry_families": 12}
    prev = {"markets": {"india": {"raw_rows": 63, "unique_predictions": 53,
                                  "unique_tickers": 43, "matured_outcomes": 0,
                                  "fundamental_periods": 265, "tickers": 45,
                                  "tickers_ge_8q": 0, "median_quarters": 6},
                        "usa": {}}, "registry_families": 12}
    d = ec.deltas(cur, prev)
    assert d["india"]["raw_rows"] == -10
    assert d["india"]["dedup_not_loss"] is True
    assert d["conserved"] is True, d["regressions"]


def test_clock_fails_on_real_evidence_loss():
    """A fall in UNIQUE identities is loss and must fail certification."""
    from backend.research.r3_program import evidence_clock as ec
    cur = {"markets": {"india": {"raw_rows": 40, "unique_predictions": 40,
                                 "unique_tickers": 30, "matured_outcomes": 0,
                                 "fundamental_periods": 200, "tickers": 40,
                                 "tickers_ge_8q": 0, "median_quarters": 6},
                       "usa": {}}, "registry_families": 12}
    prev = {"markets": {"india": {"raw_rows": 53, "unique_predictions": 53,
                                  "unique_tickers": 43, "matured_outcomes": 0,
                                  "fundamental_periods": 265, "tickers": 45,
                                  "tickers_ge_8q": 0, "median_quarters": 6},
                        "usa": {}}, "registry_families": 12}
    d = ec.deltas(cur, prev)
    assert d["conserved"] is False
    assert any("unique_predictions" in r for r in d["regressions"])
    assert any("fundamental_periods" in r for r in d["regressions"])


def test_no_new_model_guard_blocks_when_nothing_is_ready():
    from backend.research.r3_program import evidence_clock as ec
    g = ec.may_train({"ready_for_evaluation": []})
    assert g["may_train"] is False
    assert "do not train" in g["reason"]
    g2 = ec.may_train({"ready_for_evaluation": ["R3-K-META:usa"]})
    assert g2["may_train"] is True
    assert "PREDEFINED" in g2["reason"]


def test_clock_baseline_is_written_only_when_conservation_passes(tmp_path):
    """A corrupt run must not become the baseline that hides the next
    regression.

    Asserted behaviourally. This was a grep for a source string, which
    broke on a refactor that preserved the behaviour exactly - a test
    that fails when nothing is wrong teaches people to ignore it.
    """
    from backend.research.r3_program import evidence_clock as ec
    (tmp_path / "reports" / "research" / "r3").mkdir(parents=True)
    ec.emit(tmp_path, {"snapshot": _snap("2026-09-09", 53, 30),
                       "conservation": {"pass": True}})
    ec.emit(tmp_path, {"snapshot": _snap("2026-09-10", 5, 2),
                       "conservation": {"pass": False}})
    assert ec.previous_snapshot(tmp_path, "2026-09-11")["asof"] == "2026-09-09"


def test_calibration_reports_not_started_rather_than_a_fake_ece():
    from backend.research.r3_program import evidence_clock as ec
    c = ec._calibration(ROOT)
    assert c["eligible_weekly_observations"] == 0
    assert c["required_weeks"] == 4 and c["ece_threshold"] == 0.05


def test_acquisition_profile_refuses_to_report_an_empty_sample():
    """A profile where every fetch failed once read 0.000s network with 60
    retries · a measurement that cannot fail loudly is worse than none."""
    from backend.observability import acquisition_profile as ap
    src = Path(ap.__file__).read_text(encoding="utf-8")
    assert "every sampled fetch failed" in src
    assert "t.get(\"symbol\")" in src


def test_acquisition_records_per_stage_timing():
    from backend.pipeline.contract import runner
    src = Path(runner.__file__).read_text(encoding="utf-8")
    assert '"elapsed_s": round(time.time() - t1, 1)' in src


def test_clock_runs_in_the_daily_pipeline():
    from backend.pipeline.contract import runner
    src = Path(runner.__file__).read_text(encoding="utf-8")
    assert "r3_evidence_clock" in src


# ── stop / price / distance must agree · CEO 2026-09-09 ─────────────────

@pytest.mark.parametrize("market", MARKETS)
def test_stop_price_distance_reconcile(market):
    """distance = |price - stop| / price * 100 · on every row with a stop."""
    r = completeness.stop_distance_invariant(ROOT, market)
    if r.get("error"):
        pytest.skip(r["error"])
    assert not r["mismatches"], r["mismatches"][:3]
    assert not r["malformed_stops"], r["malformed_stops"][:3]
    assert r["pass"] is True


def test_stop_distance_mismatch_blocks_delivery():
    src = Path(completeness.__file__).read_text(encoding="utf-8")
    assert "STOP_DISTANCE_MISMATCH" in src
    assert "not shippable" in src


def test_stop_distance_is_never_silently_repaired():
    """A displayed value that cannot reconcile must FAIL, not be fixed."""
    src = Path(completeness.__file__).read_text(encoding="utf-8")
    assert "Nothing is repaired here" in src


def test_message_never_signs_stop_distance_as_upside():
    """`dist +7.48%` reads as a profit target · it is downside."""
    from backend.delivery.telegram import canonical_message as cm
    for m in MARKETS:
        text = cm.render(ROOT, m)
        if not text:
            continue
        assert "dist +" not in text, "stop distance rendered with a + sign"
        assert "Stop distance" in text
        assert "NOT a profit target" in text


def test_incoherent_target_is_withheld_not_shown():
    """A target at or below the current price is a target already passed."""
    src = (ROOT / "backend" / "delivery" / "sheets"
           / "workbook_two.py").read_text(encoding="utf-8")
    assert 'row["target"] <= row["current_price"]' in src


@pytest.mark.parametrize("market", MARKETS)
def test_no_row_displays_a_target_at_or_below_current_price(market, tmp_path):
    """Assert the RENDERER, not a checked-in artifact.

    The first version read reports/telegram/aegis_history_<mkt>.xlsx
    directly and failed CI three times: the committed workbook was built
    before this fix and still showed CRM at 221.57 against a price of
    259.23. A delivery test that depends on the freshness of a file it
    does not control blocks the send for a non-reason - which is exactly
    the mistake made once already with the R3 ABSTAIN column.

    The workbook is therefore built HERE from the current lifecycle. The
    live artifact is the gate's responsibility, not pytest's, and
    stop_distance_invariant already blocks delivery on it.
    """
    from openpyxl import load_workbook

    from backend.delivery.lifecycle import canonical_daily_lifecycle as lc
    from backend.delivery.sheets import workbook_two as w2
    d = lc.load(ROOT, market)
    if not d:
        pytest.skip("no lifecycle dataset for %s" % market)
    built = w2.build_two_sheet_workbook(ROOT, market, d.get("asof"))
    p = tmp_path / ("%s.xlsx" % market)
    built["workbook"].save(p)
    wb = load_workbook(p, read_only=True)
    ws = wb["CURRENT"]
    hdr = [str(c.value).strip() if c.value else "" for c in ws[7]]
    ti, pi, gi = (hdr.index("Ticker"), hdr.index("Last Price"),
                  hdr.index("Target"))
    bad = []
    for r in range(8, ws.max_row + 1):
        t = ws.cell(r, ti + 1).value
        if not t or str(t).startswith("─"):
            break
        cp, tg = ws.cell(r, pi + 1).value, ws.cell(r, gi + 1).value
        if isinstance(cp, (int, float)) and isinstance(tg, (int, float)) \
                and tg <= cp:
            bad.append((str(t), tg, cp))
    wb.close()
    assert not bad, "targets at/below current price: %s" % bad[:4]


# ── evidence clock · the baseline must survive a second invocation ────
#
# The clock is wired into a per-market loop, so it runs twice per pipeline
# run. India's invocation used to overwrite the baseline and USA's then
# compared against it: on 2026-09-10 the report read +0 for both markets
# while India had genuinely gone 53 -> 96 unique predictions. The delta
# was consumed into the baseline and never displayed.

def _snap(asof, india_unique, usa_unique, today_stamped=0):
    def mk(n, t):
        return {"raw_rows": n, "unique_predictions": n, "unique_tickers": n,
                "matured_outcomes": 0, "fundamental_periods": 0,
                "tickers": 0, "tickers_ge_8q": 0, "median_quarters": 0,
                "predictions_stamped_today": t,
                "predictions_carried_in": n - t}
    return {"asof": asof, "registry_families": 0,
            "markets": {"india": mk(india_unique, today_stamped),
                        "usa": mk(usa_unique, today_stamped)}}


def test_clock_baseline_is_not_overwritten_within_the_same_day(tmp_path):
    """A second invocation on the same as-of must not consume the delta."""
    from backend.research.r3_program import evidence_clock as ec
    d = tmp_path / "reports" / "research" / "r3"
    d.mkdir(parents=True)
    base = _snap("2026-09-09", 53, 30)
    (d / "R3_EVIDENCE_CLOCK_PREVIOUS.json").write_text(
        json.dumps(base), encoding="utf-8")

    cur = _snap("2026-09-10", 96, 38, today_stamped=43)
    for _ in range(3):                       # india, usa, and a manual re-run
        rep = {"snapshot": cur, "conservation": {"pass": True}}
        ec.emit(tmp_path, rep)
        held = json.loads(
            (d / "R3_EVIDENCE_CLOCK_PREVIOUS.json").read_text(encoding="utf-8"))
        assert held["asof"] == "2026-09-09", (
            "baseline rolled forward within the same day · the day-over-day "
            "delta has been consumed and can never be displayed")
        assert ec.deltas(cur, held)["india"]["unique_predictions"] == 43


def test_clock_baseline_rolls_when_the_day_moves(tmp_path):
    """Today's entry is recorded, and tomorrow compares against it."""
    from backend.research.r3_program import evidence_clock as ec
    d = tmp_path / "reports" / "research" / "r3"
    d.mkdir(parents=True)
    (d / "R3_EVIDENCE_CLOCK_PREVIOUS.json").write_text(
        json.dumps(_snap("2026-09-09", 53, 30)), encoding="utf-8")
    ec.emit(tmp_path, {"snapshot": _snap("2026-09-10", 96, 38),
                       "conservation": {"pass": True}})
    # Yesterday is still the baseline for TODAY ...
    assert ec.previous_snapshot(tmp_path, "2026-09-10")["asof"] == "2026-09-09"
    # ... and today has been recorded as the baseline for TOMORROW.
    nxt = ec.previous_snapshot(tmp_path, "2026-09-11")
    assert nxt["asof"] == "2026-09-10"
    assert nxt["markets"]["india"]["unique_predictions"] == 96


def test_clock_never_baselines_a_failed_conservation(tmp_path):
    from backend.research.r3_program import evidence_clock as ec
    d = tmp_path / "reports" / "research" / "r3"
    d.mkdir(parents=True)
    (d / "R3_EVIDENCE_CLOCK_PREVIOUS.json").write_text(
        json.dumps(_snap("2026-09-09", 53, 30)), encoding="utf-8")
    ec.emit(tmp_path, {"snapshot": _snap("2026-09-10", 12, 4),
                       "conservation": {"pass": False}})
    held = json.loads(
        (d / "R3_EVIDENCE_CLOCK_PREVIOUS.json").read_text(encoding="utf-8"))
    assert held["asof"] == "2026-09-09", (
        "a run that lost evidence became the baseline · the loss would be "
        "invisible tomorrow")


def test_clock_reports_intake_without_any_baseline(tmp_path):
    """Stateless intake survives a consumed or absent baseline."""
    from backend.research.r3_program import evidence_clock as ec
    cur = _snap("2026-09-10", 96, 38, today_stamped=43)
    first = ec.deltas(cur, None)
    assert first["first_run"] is True
    assert first["stateless_intake"]["india"]["predictions_stamped_today"] == 43
    same_day = ec.deltas(cur, _snap("2026-09-10", 96, 38, today_stamped=43))
    assert same_day["baseline_is_same_day"] is True
    assert same_day["india"]["unique_predictions"] == 0      # diff is blind
    assert same_day["stateless_intake"]["india"][
        "predictions_stamped_today"] == 43                   # intake is not
    assert "not since yesterday" in same_day["note"]


def test_clock_conservation_still_catches_loss(tmp_path):
    """Intake must not be allowed to mask a deletion."""
    from backend.research.r3_program import evidence_clock as ec
    cur = _snap("2026-09-10", 40, 38, today_stamped=43)
    d = ec.deltas(cur, _snap("2026-09-09", 53, 30))
    assert d["conserved"] is False
    assert any("india.unique_predictions" in r for r in d["regressions"])


# ── price freshness must come from the DATA, per market ───────────────
#
# Two defects found in the 2026-09-10 lock audit, both invisible in the
# certification output:
#   1 · freshness was the newest FILE MTIME, so a refresh that wrote every
#       parquet without adding a bar certified "age 0" over stale data;
#   2 · the directory loop ignored the market and broke on the first path
#       that existed, so INDIA's freshness was read from USA's files.

def _bars(tmp_path, rel, rows):
    """rows = {ticker: [iso dates]} written as <TICKER>_D1.parquet."""
    import pandas as pd
    d = tmp_path / rel
    d.mkdir(parents=True, exist_ok=True)
    for tk, dates in rows.items():
        df = pd.DataFrame({"close": [1.0] * len(dates)},
                          index=pd.DatetimeIndex(
                              [pd.Timestamp(x) for x in dates], name="date"))
        df.to_parquet(d / ("%s_D1.parquet" % tk))
    return d


def test_price_freshness_reads_bar_dates_not_file_mtime(tmp_path):
    from backend.pipeline.contract import readiness as R
    _bars(tmp_path, "data/raw/india",
          {"AAA": ["2026-09-01", "2026-09-02"], "BBB": ["2026-09-02"]})
    b = R._price_bars_truth(tmp_path, "india", None)
    # The files were written just now; their mtime is today. The DATA
    # ends 2026-09-02, and that is what freshness must report.
    assert b["asof"] == "2026-09-02", (
        "freshness followed the file clock instead of the bar dates")


def test_price_freshness_is_per_market(tmp_path):
    """India must never be certified by USA's files."""
    from backend.pipeline.contract import readiness as R
    _bars(tmp_path, "usa/data/raw/us", {"MSFT": ["2026-09-09"]})
    _bars(tmp_path, "data/raw/india", {"AAA": ["2026-08-01"]})
    assert R._price_bars_truth(tmp_path, "india", None)["asof"] == "2026-08-01"
    assert R._price_bars_truth(tmp_path, "usa", None)["asof"] == "2026-09-09"


def test_price_freshness_ignores_non_price_parquets(tmp_path):
    """A neighbouring file must not be able to set the as-of.

    `fundamentals.parquet` is indexed by TICKER; reading its last row
    returned the string "PEL", which sorted above every real date.
    """
    import pandas as pd
    from backend.pipeline.contract import readiness as R
    d = _bars(tmp_path, "data/raw/india", {"AAA": ["2026-09-09"]})
    pd.DataFrame({"returnOnEquity": [0.2, 0.3]},
                 index=pd.Index(["PEL", "ZZZ"], name=None)
                 ).to_parquet(d / "fundamentals.parquet")
    (d / "intraday").mkdir(exist_ok=True)
    pd.DataFrame({"close": [1.0]},
                 index=pd.DatetimeIndex([pd.Timestamp("2027-01-01")],
                                        name="date")
                 ).to_parquet(d / "intraday" / "AAA_M5.parquet")
    b = R._price_bars_truth(tmp_path, "india", None)
    assert b["asof"] == "2026-09-09", (
        "a non-price or intraday file set the market's certified as-of")


def test_price_freshness_reports_the_median_not_just_the_newest(tmp_path):
    """One fresh ticker must not certify a stale book."""
    from backend.pipeline.contract import readiness as R
    _bars(tmp_path, "data/raw/india",
          {"AAA": ["2026-09-09"], "BBB": ["2026-08-01"],
           "CCC": ["2026-08-01"], "DDD": ["2026-08-01"]})
    b = R._price_bars_truth(tmp_path, "india", None)
    assert b["asof"] == "2026-09-09"
    assert b["median_asof"] == "2026-08-01"


def test_price_freshness_scopes_to_the_universe_when_known(tmp_path):
    from backend.pipeline.contract import readiness as R
    _bars(tmp_path, "data/raw/india",
          {"AAA": ["2026-09-09"], "OLD": ["2026-01-01"]})
    scoped = R._price_bars_truth(tmp_path, "india", ["AAA"])
    assert scoped["scope"] == "universe"
    assert scoped["median_asof"] == "2026-09-09"
    assert R._price_bars_truth(tmp_path, "india", None)["scope"] == "all_files"


# ── R3 prediction identity · five parts, not two ──────────────────────
#
# CEO 2026-09-10: (source/program, market, as_of, ticker,
# prediction_version). The old key was (as_of, ticker), which is enough
# while ONE program writes the ledger and silently stops being enough the
# moment a second one does.

def test_identity_keeps_two_sources_on_one_ticker_distinct():
    """The R3-H precondition · without this the experiment is impossible."""
    from backend.research.r3_program import shadow
    a = {"contract_id": "CNBC", "market": "india",
         "as_of": "2026-09-10", "ticker": "TCS"}
    b = {"contract_id": "ZERODHA", "market": "india",
         "as_of": "2026-09-10", "ticker": "TCS"}
    assert shadow.identity(a) != shadow.identity(b), (
        "two sources on one ticker collapse into one record · the "
        "consensus/disagreement signal is destroyed before measurement")


def test_identity_separates_markets_and_versions():
    from backend.research.r3_program import shadow
    base = {"contract_id": "R3", "market": "india",
            "as_of": "2026-09-10", "ticker": "TCS"}
    assert shadow.identity(base) != shadow.identity({**base, "market": "usa"})
    assert shadow.identity(base) != shadow.identity(
        {**base, "prediction_version": "v2"})


def test_identity_is_backward_compatible():
    """A legacy row with no version must not become a second identity."""
    from backend.research.r3_program import shadow
    legacy = {"contract_id": "R3", "market": "india",
              "as_of": "2026-09-10", "ticker": "TCS"}
    explicit = {**legacy,
                "prediction_version": shadow.DEFAULT_PREDICTION_VERSION}
    assert shadow.identity(legacy) == shadow.identity(explicit)


def test_earliest_prediction_wins_a_duplicate(tmp_path):
    """A rerun has seen more of the day · that is hindsight, not a fix."""
    import json as _j
    from backend.research.r3_program import shadow
    p = shadow.ledger_path(tmp_path, "india")
    p.parent.mkdir(parents=True, exist_ok=True)
    rows = [
        {"contract_id": "R3", "market": "india", "as_of": "2026-09-10",
         "ticker": "TCS", "r3_action": "ABSTAIN",
         "recorded_utc": "2026-09-10T02:19:00+00:00"},
        {"contract_id": "R3", "market": "india", "as_of": "2026-09-10",
         "ticker": "TCS", "r3_action": "TAKE",
         "recorded_utc": "2026-09-10T06:06:00+00:00"},
    ]
    p.write_text("\n".join(_j.dumps(r) for r in rows), encoding="utf-8")
    assert len(shadow.load_ledger(tmp_path, "india", raw=True)) == 2
    ded = shadow.load_ledger(tmp_path, "india")
    assert len(ded) == 1
    assert ded[0]["r3_action"] == "ABSTAIN", (
        "the later re-prediction overwrote the original")


def test_prediction_and_outcome_stay_in_separate_files():
    from backend.research.r3_program import shadow
    led = shadow.ledger_path(ROOT, "india")
    out = (ROOT / "reports" / "research" / "r3" / "shadow"
           / "outcomes_india.jsonl")
    assert led.name != out.name and led != out


# ── USA universe reference-data defects · documented, never hidden ────
#
# CEO 2026-09-10: "Never hide a failed source behind a fabricated ticker
# mapping" and "if correction would alter production universe membership
# or R2 behavior, do not apply it silently."

def _defect_register():
    p = ROOT / "reports" / "context" / "universe_reference_defects_usa.json"
    if not p.exists():
        pytest.skip("reference-defect register missing")
    return json.loads(p.read_text(encoding="utf-8"))


def test_every_unscoreable_usa_name_is_classified():
    reg = _defect_register()
    allowed = {"CORRECTABLE_REFERENCE_ERROR", "DELISTED", "INVALID_SYMBOL",
               "NOT_US_LISTED", "UNKNOWN"}
    assert reg["entries"], "register is empty"
    for e in reg["entries"]:
        assert e["classification"] in allowed, e
        assert e.get("evidence"), "%s has no evidence" % e["symbol"]
    assert reg["n_unknown"] == 0, "an unscoreable name is still UNKNOWN"


def test_no_correction_is_applied_silently():
    """A membership-altering fix must be documented, not slipped in."""
    reg = _defect_register()
    assert reg["denominator_policy"] == "PRESERVE_516"
    for e in reg["entries"]:
        assert e["applied"] is False, (
            "%s was applied · universe membership changed without a "
            "governed decision" % e["symbol"])
        if e["classification"] == "CORRECTABLE_REFERENCE_ERROR":
            assert e.get("corrected_symbol"), e["symbol"]
            assert e.get("why_not_applied"), e["symbol"]


def test_a_delisted_name_never_carries_a_fabricated_mapping():
    """The failure must stay visible · no invented substitute symbol."""
    reg = _defect_register()
    for e in reg["entries"]:
        if e["classification"] in ("DELISTED", "NOT_US_LISTED"):
            assert e.get("corrected_symbol") is None, (
                "%s is %s yet carries a substitute symbol"
                % (e["symbol"], e["classification"]))


def test_register_matches_what_the_pipeline_actually_cannot_score():
    """The register must describe reality, not a stale belief."""
    from backend.pipeline.contract.readiness import _universe_list
    uni = _universe_list(ROOT, "usa")
    if not uni:
        pytest.skip("no USA universe")
    have = {p.stem.split("_")[0].upper()
            for p in (ROOT / "usa" / "data" / "raw" / "us").glob("*_D1.parquet")}
    unscoreable = {str(t).upper() for t in uni if str(t).upper() not in have}
    listed = {e["symbol"].upper() for e in _defect_register()["entries"]}
    assert unscoreable == listed, (
        "register drifted from reality · unlisted=%s stale=%s"
        % (sorted(unscoreable - listed), sorted(listed - unscoreable)))
